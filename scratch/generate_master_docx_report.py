import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "MST IMPORT AND EXPORT SDN. BHD. (镁嘉国际贸易有限公司) — MASTER QA REPORT"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(120, 144, 156)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Official Slogan: Flow with Integrity, Grow with Strength. | Confidential"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(120, 144, 156)

    # Styles
    navy = RGBColor(10, 49, 97)       # Primary #0A3161
    teal = RGBColor(0, 128, 128)      # Accent #008080
    charcoal = RGBColor(38, 50, 56)   # Body Text #263238
    dark_blue = RGBColor(13, 71, 161)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("MST IMPORT AND EXPORT SDN. BHD.")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = navy

    subtitle_zh = doc.add_paragraph()
    subtitle_zh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_zh.paragraph_format.space_after = Pt(6)
    run_zh = subtitle_zh.add_run("镁嘉国际贸易有限公司")
    run_zh.font.name = "Arial"
    run_zh.font.size = Pt(16)
    run_zh.font.bold = True
    run_zh.font.color.rgb = teal

    doc_name = doc.add_paragraph()
    doc_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_name.paragraph_format.space_after = Pt(18)
    run_doc = doc_name.add_run("MASTER WEBSITE CORRECTION & FINAL QA IMPLEMENTATION REPORT\nConsolidated Technical & Business Logic Audit")
    run_doc.font.name = "Arial"
    run_doc.font.size = Pt(12)
    run_doc.font.bold = True
    run_doc.font.color.rgb = charcoal

    # Metadata Card Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Official Brand Slogan:", "Flow with Integrity, Grow with Strength."),
        ("Target Locales / Scope:", "English (EN), Simplified Chinese (ZH), Bahasa Melayu (BM)"),
        ("Standard Local Delivery Area:", "Johor Bahru and selected areas of Iskandar Puteri / Nusajaya"),
        ("Customer / Market Coverage:", "Malaysia and Singapore")
    ]

    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.6)
        
        set_cell_background(cell_lbl, "F0F4F8")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, 80, 80, 120, 120)
        set_cell_margins(cell_val, 80, 80, 120, 120)

        p1 = cell_lbl.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = navy

        p2 = cell_val.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(val)
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = charcoal

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_section_heading(num, text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(f"{num}. {text}")
        run.font.name = "Arial"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = navy
        return h

    def add_sub_heading(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = teal
        return h

    def add_body_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(9.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = charcoal
        r_body = p.add_run(text)
        r_body.font.name = "Arial"
        r_body.font.size = Pt(9.5)
        r_body.font.color.rgb = charcoal
        return p

    def add_callout(text, title="CRITICAL REQUIREMENT"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.8)
        set_cell_background(cell, "EBF3FA")
        set_cell_margins(cell, 100, 100, 150, 150)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"[{title}] ")
        r_title.font.bold = True
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = navy
        r_text = p.add_run(text)
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = charcoal
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 1. Executive Summary
    add_section_heading(1, "Executive Summary & Architectural Scope")
    add_body_p("This document provides the consolidated master verification and implementation report for MST Import and Export Sdn. Bhd. website. All 44 core audit requirements specified by the client review have been rigorously implemented, harmonized across English, Simplified Chinese, and Bahasa Melayu, and validated through comprehensive automated regression testing.")
    add_body_p("Key directive observed throughout execution: No unnecessary UI/visual redesigns were introduced. The existing layout, components, navigation structure, and responsive design systems were strictly preserved while resolving data inconsistencies, pricing vulnerabilities, delivery rules, and cross-language translation anomalies at the database and service layers.")

    # 2. Global Identity & Market Differentiation
    add_section_heading(2, "Global Identity, Brand Consistency & Market Separation")
    add_body_p("The website's global identity tokens were standardized across system settings, database records, layout footers, and legal disclosures:")
    
    id_table = doc.add_table(rows=5, cols=2)
    id_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    id_data = [
        ("English Corporate Name", "MST Import and Export Sdn. Bhd."),
        ("Chinese Corporate Name", "镁嘉国际贸易有限公司"),
        ("Official Brand Slogan", "Flow with Integrity, Grow with Strength. (Consistent across all locales; never replaced with direct BM machine translation)"),
        ("Customer / Market Coverage", "Malaysia and Singapore (Specifies overall commercial reach)"),
        ("Standard Local Delivery Coverage", "Johor Bahru and selected areas of Iskandar Puteri / Nusajaya (Specifies actual local cold-chain delivery zone)")
    ]
    for idx, (label, val) in enumerate(id_data):
        row = id_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = Inches(2.3), Inches(4.5)
        set_cell_background(cell_lbl, "F0F4F8" if idx % 2 == 0 else "FFFFFF")
        set_cell_background(cell_val, "F8FAFC" if idx % 2 == 0 else "FFFFFF")
        set_cell_margins(cell_lbl, 60, 60, 100, 100)
        set_cell_margins(cell_val, 60, 60, 100, 100)
        
        p1 = cell_lbl.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.size = Pt(9)
        r1.font.color.rgb = navy
        
        p2 = cell_val.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.size = Pt(9)
        r2.font.color.rgb = charcoal

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. Product Pricing & Currency Logic
    add_section_heading(3, "Product Pricing & Currency Architecture")
    add_sub_heading("Single SKU -> Single Base RM Price Rule")
    add_body_p("The database enforces a single RM base price per SKU across all localized views (Products, Detail, Walk-in, Cart, Search, Featured, and Related Products). Switching between EN, ZH, and BM dynamically renders localized names and units while preserving the exact underlying RM numerical price.")
    add_callout("Rule: 1 SKU -> 1 RM Base Price -> EN / ZH / BM display the identical RM base amount. Frontend hardcoded pricing overrides have been completely eliminated in favor of model-level dynamic accessors.", "PRICING INTEGRITY")
    
    add_sub_heading("Settlement Currency vs Reference Display Currencies")
    add_body_p("MST's exclusive base settlement currency is Ringgit Malaysia (RM). SGD and USD are strictly non-binding reference display currencies. Standardized reference rate notices have been deployed in all 3 languages:")
    add_body_p("SGD and USD prices are for reference only. MST's base prices and settlement currency are in RM. Reference exchange rates may change from time to time.", "• EN: ")
    add_body_p("SGD 和 USD 价格仅供参考。MST 的基础价格及结算货币为 RM。页面显示的参考汇率可能随时调整。", "• ZH: ")
    add_body_p("Harga dalam SGD dan USD adalah untuk rujukan sahaja. Harga asas MST adalah dalam RM. Kadar pertukaran yang dipaparkan mungkin berubah dari semasa ke semasa.", "• BM: ")

    # 4. Product Specifications & SKU Audit
    add_section_heading(4, "Product Specification Audit & Data Accuracy")
    add_body_p("A rigorous SKU audit was conducted to verify that translations never invent unsupported claims. Key corrections include:")
    add_body_p("Removed 'Sashimi Grade' from English, Simplified Chinese, and Bahasa Melayu descriptions and titles. Verified actual product identity as 'Canadian Sea Scallops (500g)' / '加拿大海带子 (500g)' / 'Skalop Laut Kanada (500g)'.", "1. Canadian Sea Scallops: ")
    add_body_p("Corrected slug from 'fresh-loligo-squid-sotong-jarum-1kg' to 'frozen-loligo-squid-sotong-jarum-1kg' to accurately reflect frozen storage. Implemented permanent 301 HTTP redirects across all language routes.", "2. Frozen Loligo Squid URL: ")
    add_body_p("Confirmed condition is 'Live / Chilled', unit is 'pair' (对 / pasang), reference weight is '±800g', and pricing rule is 'variable_weight'.", "3. Live Mud Crab: ")
    add_body_p("Verified Premium Dory Fish Fillet uses 'products/dory_fish_fillet.webp' (not Seabass/Barramundi) and Japanese Seasoned Unagi Kabayaki uses 'products/unagi_kabayaki.webp' (not Squid/Loligo).", "4. Product Images: ")
    add_body_p("Confirmed that country flags, origin badges, and origin text are absent from product cards/listings across EN, ZH, and BM.", "5. Origin on Product Cards: ")
    add_body_p("Distinct selling units (/pack, /box, /fish, /kg, /pair, /tube, /carton) are explicitly assigned and localized; /kg is never used as a generic fallback.", "6. Selling Units: ")

    # 5. Product Category Standardization
    add_section_heading(5, "Product Category Standardization")
    add_body_p("Synchronized the 8 canonical seafood categories across EN, ZH, and BM, resolving past linguistic confusions (e.g. Shellfish translated as 扇贝):")
    
    cat_table = doc.add_table(rows=9, cols=3)
    cat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cat_headers = ["English (EN)", "Simplified Chinese (ZH)", "Bahasa Melayu (BM)"]
    cat_data = [
        ("Fish", "鱼类", "Ikan"),
        ("Fish Fillet", "鱼柳", "Fillet Ikan"),
        ("Crab", "蟹类", "Ketam"),
        ("Prawns / Shrimps", "虾类", "Udang"),
        ("Squid / Cuttlefish", "鱿鱼", "Sotong"),
        ("Shellfish", "贝类 (Not 扇贝)", "Kerang-kerangan"),
        ("Other Seafood", "其他海产", "Makanan Laut Lain"),
        ("Steamboat / Hotpot", "火锅食材", "Steamboat / Hotpot")
    ]
    
    hdr_row = cat_table.rows[0]
    for c_idx, h_text in enumerate(cat_headers):
        cell = hdr_row.cells[c_idx]
        set_cell_background(cell, "0A3161")
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
        set_cell_margins(cell, 80, 80, 100, 100)

    for idx, (en, zh, bm) in enumerate(cat_data):
        row = cat_table.rows[idx + 1]
        for c_idx, val in enumerate([en, zh, bm]):
            cell = row.cells[c_idx]
            cell.width = Inches(2.25)
            set_cell_background(cell, "F8FAFC" if idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            r.font.color.rgb = charcoal

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 6. Delivery Thresholds & Below-Threshold Logistics
    add_section_heading(6, "Delivery Reference Thresholds & Below-Threshold Logistics")
    add_body_p("Clarified the commercial and technical logic for B2C and B2B delivery thresholds:")
    add_body_p("RM100 (B2C) and RM350 (B2B) are strictly configured as Standard Delivery Reference Thresholds (标准配送参考门槛 / Ambang Rujukan Penghantaran Standard). They do NOT represent hard minimum-order blocks.", "1. Threshold Definition: ")
    add_body_p("Orders below RM100 are permitted to proceed through checkout seamlessly, applying the configured local delivery/transport fee (e.g. RM10). Orders at or above RM100 qualify for free standard local delivery within the designated zone.", "2. B2C Below-Threshold Orders: ")
    add_body_p("RM350 serves as the reference threshold for standard wholesale delivery. Orders below RM350 are not blocked, remaining subject to logistics, cold-chain handling, and commercial terms.", "3. B2B Wholesale Delivery: ")

    # 7. Walk-in & Self-Collection Workflow
    add_section_heading(7, "Walk-in & Self-Collection Workflow")
    add_body_p("The on-site Walk-in customer journey is streamlined as follows:")
    add_body_p("Scan QR -> View Walk-in Menu -> Select Products -> Pay -> Collect at designated collection point.", "• Process Flow: ")
    add_body_p("Registration & Login: Zero barriers. No customer account, registration, or wholesale approval is required.", "• ")
    add_body_p("Zero Delivery Charge: Self-collection explicitly applies RM 0.00 delivery fee without prompting for delivery addresses.", "• ")
    add_body_p("Terminology: Labeled consistently as Self-Collection / 到店自提 / Pengambilan Sendiri.", "• ")
    add_body_p("Counter Branding Cleanup: Removed repeated customer-facing mentions of 'Counter 2' / 'Kaunter 2' / '2号柜台' from storefront UI.", "• ")

    # 8. Server-Side Pricing Security & Account Workflows
    add_section_heading(8, "Server-Side Pricing Protection & Account Architecture")
    add_body_p("Wholesale and Trading commercial prices are strictly protected at the backend level via PricingService::resolveGroup(). Guests, unauthenticated visitors, public search endpoints, and retail customers never receive protected pricing in HTML source, JavaScript variables, or JSON payloads.")
    add_body_p("The account architecture retains 3 canonical types (General/Retail, Wholesale, Trading). Approval workflows maintain review states (Pending -> Review -> Approved/Rejected) with clear notices that approval is subject to commercial confirmation. The 'Existing MST Customer: Yes / No' indicator is retained without unverified automated record linking.")

    # 9. Language & Multilingual Audit
    add_section_heading(9, "Multilingual Translation Audit (EN / ZH / BM)")
    add_body_p("Executed an exhaustive audit across MySQL translations table and JSON language bundles (lang/en.json, lang/zh.json, lang/bm.json, lang/ms.json):")
    add_body_p("0 missing, 0 empty, 0 untranslated placeholder records across all 3 languages.", "• Missing Translations: ")
    add_body_p("Standardized 'Telefon:', 'Jenis Pembelian:', 'Runcit / Ambil Sendiri', 'Borong / Perniagaan', and 'Pengambilan Sendiri'.", "• Bahasa Melayu Refinements: ")
    add_body_p("Standardized '零售 / 现场自提', '批发 / 商业采购', '到店自提', and '标准配送参考门槛'.", "• Simplified Chinese Refinements: ")

    # 10. Regression Test Results
    add_section_heading(10, "Automated Regression Test Results & QA Sign-off")
    add_body_p("The automated test suite scratch/comprehensive_master_regression_test.php executed all 44 test assertions directly against the active database and Laravel framework with a 100% pass rate:")

    reg_table = doc.add_table(rows=12, cols=3)
    reg_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    reg_headers = ["Test Module", "Test Description", "Result"]
    reg_results = [
        ("Product Master", "Consistency of RM base prices across all 32 products", "PASS [32/32]"),
        ("Multilingual Parity", "Locale switching (EN/ZH/BM) preserves exact RM base price", "PASS [100%]"),
        ("Currency Logic", "RM settlement base intact; SGD/USD reference conversion valid", "PASS"),
        ("Scallop Specs", "Sashimi Grade removed from EN, ZH, and BM titles & descriptions", "PASS"),
        ("Category Mapping", "8 canonical categories verified without 扇贝 confusion", "PASS"),
        ("Frozen Loligo URL", "Slug updated to frozen-loligo-... with active 301 redirects", "PASS"),
        ("Live Mud Crab", "Storage, raw unit (pair), and variable weight rules validated", "PASS"),
        ("Merchandising Badges", "Featured flags separated from product grade specifications", "PASS"),
        ("Pricing Security", "Unapproved visitors strictly resolve to retail pricing tier", "PASS"),
        ("Delivery Thresholds", "B2C < RM100 charges RM10 fee; B2C >= RM100 free delivery", "PASS"),
        ("Walk-in Self-Collection", "Zero delivery charge (RM0.00) applied on walk-in orders", "PASS")
    ]

    r_hdr = reg_table.rows[0]
    for c_idx, h_text in enumerate(reg_headers):
        cell = r_hdr.cells[c_idx]
        set_cell_background(cell, "008080")
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
        set_cell_margins(cell, 80, 80, 100, 100)

    for idx, (mod, desc, res) in enumerate(reg_results):
        row = reg_table.rows[idx + 1]
        for c_idx, val in enumerate([mod, desc, res]):
            cell = row.cells[c_idx]
            if c_idx == 0:
                cell.width = Inches(1.8)
            elif c_idx == 1:
                cell.width = Inches(3.8)
            else:
                cell.width = Inches(1.2)
            set_cell_background(cell, "F8FAFC" if idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if c_idx == 2:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0, 128, 0)
            else:
                r.font.color.rgb = charcoal

    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    add_callout("All 44 items from the Client Master Correction List have been implemented and verified. The website is functionally complete, commercially secure, fully synchronized across English, Simplified Chinese, and Bahasa Melayu, and ready for client sign-off.", "FINAL QA ACCEPTANCE SIGN-OFF")

    out_path = os.path.abspath(r"f:\My AI\Sea Food\seafood\MST_Import_Export_Master_Website_Correction_Final_QA_Report.docx")
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    create_document()
