import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_manual():
    doc = docx.Document()

    # Set page margins: 0.75 in top/bottom, 0.75 in left/right
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Color Palette
    c_navy      = RGBColor(15, 39, 74)     # #0F274A (Primary Brand)
    c_blue      = RGBColor(29, 78, 216)    # #1D4ED8 (Accent Blue)
    c_teal      = RGBColor(15, 118, 110)   # #0F766E (Ocean Teal)
    c_slate     = RGBColor(71, 85, 105)    # #475569 (Secondary Text)
    c_charcoal  = RGBColor(30, 41, 59)     # #1E293B (Body Text)
    c_dark      = RGBColor(15, 23, 42)     # #0F172A (Headings)
    c_green     = RGBColor(22, 101, 52)    # #166534 (Success / Confirmed)
    c_amber     = RGBColor(180, 83, 9)     # #B45309 (Warning / Notice)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}" w:color="auto" w:val="clear"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_table_border(table, color="CBD5E1"):
        tbl = table._tbl
        tblPr = tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement('w:tblPr')
            tbl.insert(0, tblPr)
        tblBorders = OxmlElement('w:tblBorders')
        for border_name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), color)
            tblBorders.append(border)
        tblPr.append(tblBorders)

    def add_heading(text, level=1, color=None, space_before=14, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(13.5)
            run.font.color.rgb = color or c_navy
        elif level == 2:
            run.font.size = Pt(11.5)
            run.font.color.rgb = color or c_blue
        elif level == 3:
            run.font.size = Pt(10.5)
            run.font.color.rgb = color or c_teal
        return p

    def add_body(text, size=9.5, italic=False, bold=False, color=None, space_before=2, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.italic = italic
        run.font.bold = bold
        run.font.color.rgb = color or c_charcoal
        return p

    def add_bullet(text, bold_prefix=None, size=9.5, color=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r1 = p.add_run(bold_prefix)
            r1.font.name = "Calibri"
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = c_navy
        r2 = p.add_run(text)
        r2.font.name = "Calibri"
        r2.font.size = Pt(size)
        r2.font.color.rgb = color or c_charcoal
        return p

    def add_callout(text, title=None, fill_hex="EFF6FF", border_hex="3B82F6", text_color=None):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(7.0)

        cell = table.cell(0, 0)
        set_cell_background(cell, fill_hex)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
            f'<w:top w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.line_spacing = 1.15

        if title:
            r_title = cp.add_run(title + "\n")
            r_title.font.name = "Calibri"
            r_title.font.bold = True
            r_title.font.size = Pt(10)
            r_title.font.color.rgb = c_navy

        r_text = cp.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = text_color or c_charcoal

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_image_figure(image_path, caption_text, width=Inches(6.2)):
        if os.path.exists(image_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.add_run().add_picture(image_path, width=width)

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(8)
            r_cap = p_cap.add_run(f"Figure: {caption_text}")
            r_cap.font.name = "Calibri"
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = c_slate

    # ═══════════════════════════════════════════════════════════════════════════
    # 1. HEADER / TITLE BLOCK
    # ═══════════════════════════════════════════════════════════════════════════
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.columns[0].width = Inches(7.0)
    h_cell = header_table.cell(0, 0)
    set_cell_background(h_cell, "0F274A")
    set_cell_margins(h_cell, top=160, bottom=160, left=180, right=180)

    hp = h_cell.paragraphs[0]
    hp.paragraph_format.space_before = Pt(4)
    hp.paragraph_format.space_after = Pt(2)

    r_org = hp.add_run("MST IMPORT AND EXPORT SDN. BHD. (镁嘉国际贸易有限公司)\n")
    r_org.font.name = "Calibri"
    r_org.font.size = Pt(11)
    r_org.font.bold = True
    r_org.font.color.rgb = RGBColor(147, 197, 253)

    r_title = hp.add_run("COMPLETE WEBSITE ADMIN PANEL & OPERATIONS MANUAL\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(255, 255, 255)

    r_sub = hp.add_run("Step-by-Step Practical Guide for Routine Content, Product Pricing, Translations, Delivery Zones & Order Management")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(224, 242, 254)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Metadata Grid
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_border(meta_table, "CBD5E1")
    col_widths = [Inches(3.5), Inches(3.5)]

    meta_data = [
        [("Document Purpose:", "Official Back-Office User & Operations Manual"), ("Target Audience:", "Wendy (Client Project Lead) & MST Administrators")],
        [("Platform Engine:", "Laravel Unified B2B & B2C Cold-Chain Platform"), ("System Version:", "Production 2.0 (Live LTS)")],
        [("Superadmin Account:", "admin@mst.my"), ("Date of Issuance:", "October 2026")]
    ]

    for r_idx, row in enumerate(meta_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "F8FAFC" if (r_idx % 2 == 0) else "FFFFFF")
            set_cell_margins(cell, top=65, bottom=65, left=100, right=100)
            lbl, val = meta_data[r_idx][c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(lbl + " ")
            r1.font.name = "Calibri"
            r1.font.size = Pt(8.5)
            r1.font.bold = True
            r1.font.color.rgb = c_slate
            r2 = p.add_run(val)
            r2.font.name = "Calibri"
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = c_dark

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Executive Opening Callout
    add_callout(
        "Welcome to the MST Seafood Platform Administration Manual.\n\n"
        "This handbook is designed specifically for MST management and operations staff to manage the e-commerce website independently. "
        "It provides comprehensive, step-by-step instructions on updating company contact information, managing seafood products and 4-tier pricing, "
        "modifying trilingual website translations (English, Chinese, Malay), supervising delivery zones and fees, processing customer orders, "
        "and creating 1-click database backups—all directly through your browser without requiring technical coding knowledge.",
        title="Purpose & Scope of this Operations Manual",
        fill_hex="F0FDF4",
        border_hex="16A34A"
    )

    # ═══════════════════════════════════════════════════════════════════════════
    # 2. CHAPTER 1: LOGGING IN & DASHBOARD OVERVIEW
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("1. Accessing the Admin Panel & Dashboard Navigation", level=1)
    add_body(
        "The administrative panel is protected by role-based authentication and is accessible only to verified administrators."
    )

    add_heading("1.1. How to Log In", level=2)
    add_bullet("Open your web browser (Google Chrome, Microsoft Edge, Safari, or Firefox).", bold_prefix="Step 1 — Navigate: ")
    add_bullet("Go to your website domain followed by /admin (for example: https://yourdomain.com/admin or https://yourdomain.com/login).", bold_prefix="Step 2 — URL: ")
    add_bullet("Enter your administrator email address: admin@mst.my", bold_prefix="Step 3 — Username: ")
    add_bullet("Enter your secure administrator password and click 'Sign In'.", bold_prefix="Step 4 — Password: ")
    add_bullet("Upon successful authentication, you will be redirected to the Admin Overview Dashboard.", bold_prefix="Step 5 — Dashboard: ")

    add_heading("1.2. Dashboard Interface Layout", level=2)
    add_bullet("Sidebar Navigation (Left): Provides immediate access to all management modules: Dashboard, Products, Categories, Orders, Customers, Quotations, Media Gallery, Delivery Zones, Translations, Policies, and Store Settings.", bold_prefix="• Left Navigation Menu: ")
    add_bullet("Top Bar (Header): Displays quick links to the live storefront ('View Website'), current active language, notification alerts, and your profile dropdown ('Logout').", bold_prefix="• Top Header: ")
    add_bullet("KPI Overview Cards: Summary widgets showing total sales, pending orders, wholesale approval requests, and low-stock alerts.", bold_prefix="• Summary Widgets: ")

    # ═══════════════════════════════════════════════════════════════════════════
    # 3. CHAPTER 2: STORE SETTINGS & COMPANY DETAILS
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("2. Managing Company Information, Contact Details & Operating Hours", level=1)
    add_body(
        "Routine updates to your business address, contact numbers, email addresses, and operating hours can be updated in real time via Store Settings."
    )

    add_bullet("Navigate to: Admin Panel → Store Settings (/admin/settings).", bold_prefix="Step 1: ")
    add_bullet("Click on the 'Contact Page Settings' tab.", bold_prefix="Step 2: ")
    add_bullet("Update the following fields as needed:\n"
               "  • Company Name (English): e.g., MST Import and Export Sdn. Bhd.\n"
               "  • Company Name (Chinese): e.g., 镁嘉国际贸易有限公司\n"
               "  • Physical Address: e.g., 7, Jalan SILC 2/18, Kawasan Perindustrian SILC, 79200 Iskandar Puteri, Johor\n"
               "  • Primary Phone Number: e.g., +60 13-280 0168\n"
               "  • WhatsApp Number: Enter as 601112710260 (Important: Enter numbers only, without '+' or hyphens, to ensure the direct chat link works on both mobile and desktop)\n"
               "  • Customer Service Email: info@mst.my\n"
               "  • Wholesale Inquiries Email: wholesale@mst.my\n"
               "  • Business Operating Hours: e.g., Monday – Saturday: 8:00am – 6:00pm (Sunday & Public Holidays: Closed)", bold_prefix="Step 3 — Editable Fields: ")
    add_bullet("Click the blue 'Save Settings' button at the bottom of the card. Changes take effect across your website footer, contact page, and printed invoices immediately.", bold_prefix="Step 4 — Save: ")

    # ═══════════════════════════════════════════════════════════════════════════
    # 4. CHAPTER 3: WEBSITE BANNERS, POLICIES & STATIC CONTENT
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("3. Updating Banners, Announcements & Legal Policy Pages", level=1)
    
    add_heading("3.1. Homepage Hero Banner & Meta Tags", level=2)
    add_bullet("Navigate to: Admin Panel → Store Settings → General Settings tab.", bold_prefix="Location: ")
    add_bullet("Edit the 'Site Title', 'Site Description' (shows on Google search results), and 'Meta Keywords'.", bold_prefix="Editing: ")
    add_bullet("Click 'Save Settings' to apply.", bold_prefix="Save: ")

    add_heading("3.2. Legal & Information Pages (Terms, Privacy, Refund, Shipping)", level=2)
    add_bullet("Navigate to: Admin Panel → Policies & Dynamic Pages (/admin/policies).", bold_prefix="Step 1: ")
    add_bullet("You will see a list of standard policy pages: Privacy Policy, Terms & Conditions, Refund & Return Policy, Shipping & Cold-Chain Policy, and Cookie Policy.", bold_prefix="Step 2: ")
    add_bullet("Click the 'Edit' (pencil icon) button next to the policy you wish to update.", bold_prefix="Step 3: ")
    add_bullet("Edit the text directly using the built-in rich text editor (supports headings, bold, bullet points, and links).", bold_prefix="Step 4: ")
    add_bullet("Ensure the 'Published' toggle is checked, then click 'Update Policy'.", bold_prefix="Step 5: ")

    # ═══════════════════════════════════════════════════════════════════════════
    # 5. CHAPTER 4: TRILINGUAL LOCALIZATION (EN, ZH, BM)
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("4. Managing Trilingual Content (English, Simplified Chinese, Bahasa Melayu)", level=1)
    add_body(
        "The MST platform features a complete trilingual translation management system. Content is localized in two distinct ways:"
    )

    add_heading("4.1. Localizing Product Names & Descriptions", level=2)
    add_body("When editing any product in Product Management, you have dedicated fields for each language:")
    add_bullet("English: 'Product Name' and 'Product Description'", bold_prefix="• English (EN): ")
    add_bullet("Chinese: 'Product Name (Chinese)' and 'Description (Chinese)' (e.g. 挪威三文鱼片)", bold_prefix="• Simplified Chinese (ZH): ")
    add_bullet("Malay: 'Product Name (Malay)' and 'Description (Malay)' (e.g. Filet Ikan Salmon Norway)", bold_prefix="• Bahasa Melayu (BM): ")
    add_body("When a customer toggles the language switch in the header, the product page displays the corresponding translation automatically.")

    add_heading("4.2. Localizing Website Buttons, Labels & Interface Phrases", level=2)
    add_bullet("Navigate to: Admin Panel → Multilingual Translations (/admin/translations).", bold_prefix="Step 1: ")
    add_bullet("Use the Search Box to find any button label, header title, or alert message you wish to change (e.g., search for 'Free Standard Delivery' or 'Add to Cart').", bold_prefix="Step 2: ")
    add_bullet("Edit the translation directly in the corresponding column (English, 中文, or Bahasa Melayu).", bold_prefix="Step 3: ")
    add_bullet("Click 'Save All Translations' at the bottom of the page.", bold_prefix="Step 4: ")
    add_bullet("Important Tip: If updates do not reflect immediately on the storefront, click the 'Clear Translation Cache' button at the top right of the page.", bold_prefix="Step 5 — Cache Purge: ")

    # ═══════════════════════════════════════════════════════════════════════════
    # 6. CHAPTER 5: PRODUCT CATALOG & MULTI-TIER PRICING
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("5. Product Catalog, Inventory & Multi-Tier Pricing Management", level=1)
    add_body(
        "Product Management allows full control over your seafood inventory, pricing schedules, images, and tier visibility."
    )

    add_heading("5.1. Adding a New Product", level=2)
    add_bullet("Navigate to: Admin Panel → Product Management (/admin/products) and click '+ Add New Product'.", bold_prefix="Step 1: ")
    add_bullet("Basic Info: Enter the Product Name (EN, ZH, BM), SKU (Stock Keeping Unit, e.g., SALMON-NOR-1KG), select Category, and specify Unit (e.g., kg, pack, carton, box).", bold_prefix="Step 2 — Basic Info: ")
    add_bullet("Seafood Specifications: Fill in Origin (e.g. Norway, Australia, Local Malaysia), Storage Temperature (-18°C), Glaze Percentage (e.g. 10% Glaze), and Freezing Method (IQF - Individually Quick Frozen).", bold_prefix="Step 3 — Specifications: ")
    add_bullet("Multi-Tier Pricing: Enter the distinct unit price (MYR) for each customer tier:\n"
               "  • Retail Price: Price for public guest shoppers (e.g., RM 48.00)\n"
               "  • Walk-In Price: Special promotional price for storefront QR customers (e.g., RM 44.00)\n"
               "  • Wholesale Price: Commercial discounted price for approved restaurants/caterers (e.g., RM 38.00)\n"
               "  • Trading Price: Pallet-scale rate for bulk buyers (e.g., RM 34.00)\n"
               "  • (Optional): Fixed SGD and USD price fields can be filled if you wish to override automated exchange rates.", bold_prefix="Step 4 — 4-Tier Pricing: ")
    add_bullet("Minimum Order Quantities (MOQ): Set Wholesale MOQ (e.g., 5 cartons) and Trading MOQ (e.g., 20 cartons).", bold_prefix="Step 5 — MOQs: ")
    add_bullet("Stock Tracking: Enter the current stock count and check 'Track Stock' to enable automatic stock deduction upon orders.", bold_prefix="Step 6 — Stock: ")
    add_bullet("Images: Upload a primary thumbnail image and optional gallery images (supports JPG, PNG, WebP).", bold_prefix="Step 7 — Images: ")
    add_bullet("Click 'Create Product' to publish.", bold_prefix="Step 8 — Publish: ")

    add_heading("5.2. Controlling Product Visibility Flags", level=2)
    add_bullet("Active (is_active): Check to show the product on the website; uncheck to hide it immediately.", bold_prefix="• Global Visibility: ")
    add_bullet("Walk-In Available (is_walkin_available): Check this box only if the item is physically stocked at your SILC Counter 2 storefront. If unchecked, the item is hidden from walk-in QR shoppers.", bold_prefix="• Walk-In QR Filter: ")
    add_bullet("Featured Product (is_featured): Check to display this product on the homepage 'Featured Seafood' showcase.", bold_prefix="• Homepage Showcase: ")
    add_bullet("RFQ-Only Mode (is_rfq_only): Check this box for floating-market items where you do not wish to publish a fixed price. The 'Buy' button is converted to a 'Request for Quotation' button.", bold_prefix="• Trading RFQ Mode: ")

    # Illustrated Figures
    assets_dir = os.path.join(os.path.dirname(__file__), "doc_assets")
    add_image_figure(os.path.join(assets_dir, "products_en.png"), "Live Product Catalog Display & Multi-Tier Pricing Structure")
    add_image_figure(os.path.join(assets_dir, "walkin_shop.png"), "Walk-In Counter 2 Quick-Order Interface (Filtered by is_walkin_available)")

    # ═══════════════════════════════════════════════════════════════════════════
    # 7. CHAPTER 6: DELIVERY ZONES, POSTCODES & FEES
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("6. Managing Delivery Zones, Postcodes & Fulfillment Rules", level=1)
    add_body(
        "The logistics engine determines delivery fees based on the customer's delivery postcode and order value."
    )

    add_heading("6.1. Managing Zone A Coverage & Fees", level=2)
    add_bullet("Navigate to: Admin Panel → Delivery Zones (/admin/delivery-zones).", bold_prefix="Step 1: ")
    add_bullet("Click 'Edit' on Zone A (Local JB, Iskandar Puteri & Central Suburbs).", bold_prefix="Step 2: ")
    add_bullet("Postcodes Field: Enter one 5-digit postcode per line. Zone A currently covers 79xxx, 80xxx, 81100 (Mount Austin), 81200, 81300 (Skudai), 81400 (Senai), 81550 (Gelang Patah / Setia Eco Gardens), 81750 (Masai), and 81000 (Indahpura).", bold_prefix="Step 3 — Postcodes: ")
    add_bullet("Fee Rules:\n"
               "  • Delivery Fee: Set to 0.00 (Standard fee for orders qualifying for free delivery)\n"
               "  • Below Threshold Fee: Set to 10.00 (Fee applied when order value is below RM150)\n"
               "  • Free Delivery Threshold: Set in the threshold widget (default: RM 150.00 for Retail, RM 350.00 for Wholesale)", bold_prefix="Step 4 — Fees: ")
    add_bullet("Click 'Save Zone' to apply updates.", bold_prefix="Step 5: ")

    add_heading("6.2. Special Rule for Kulai (Indahpura Only)", level=2)
    add_callout(
        "Important System Rule: Postcode 81000 covers both urban Indahpura and wider rural Kulai. "
        "The system automatically checks the delivery address:\n"
        "• If the customer enters 'Indahpura', Zone A local delivery applies.\n"
        "• If the customer enters rural Kulai, the system automatically routes the order to Outstation Cold-Chain (Zone B).\n"
        "You do not need to manually change this; the system handles it automatically.",
        title="Automated Indahpura Locality Validation",
        fill_hex="FFFBEB",
        border_hex="F59E0B"
    )

    add_heading("6.3. Outstation Orders (Zone B)", level=2)
    add_body(
        "Outstation orders are flagged with 'manual quotation required'. When a customer outside Zone A orders, "
        "the website displays: 'Outstation Transportation Fee — To Be Confirmed' and charges RM0 for shipping on Stripe. "
        "Your operations team coordinates the Styrofoam box sizing and cold-chain freight rate with the customer via WhatsApp before dispatch."
    )

    add_image_figure(os.path.join(assets_dir, "checkout_en.png"), "Storefront Checkout Interface with Automatic Postcode Validation & Zone Rules")

    # ═══════════════════════════════════════════════════════════════════════════
    # 8. CHAPTER 7: ORDERS & FULFILLMENT MANAGEMENT
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("7. Order Management, Dispatch Scheduling & Invoicing", level=1)

    add_heading("7.1. Order Processing Workflow", level=2)
    add_bullet("Navigate to: Admin Panel → Orders Dashboard (/admin/orders).", bold_prefix="Step 1: ")
    add_bullet("Filter orders by Status (Pending, Confirmed, Preparation, Shipped, Delivered, Collected), Payment Status (Paid, Unpaid), or Fulfillment Type (Delivery vs Self-Collection).", bold_prefix="Step 2 — Filter: ")
    add_bullet("Click on any Order Reference (e.g., ORD-6705F...) to open the detailed order summary.", bold_prefix="Step 3 — View Order: ")
    add_bullet("Review the customer's delivery address, special notes, selected collection time slot, and ordered seafood items.", bold_prefix="Step 4 — Review: ")
    add_bullet("Updating Status: Change the order status to 'Confirmed' once verified, 'Preparation' when packing in the cold room, and 'Shipped' or 'Collected' upon dispatch.", bold_prefix="Step 5 — Status Update: ")

    add_heading("7.2. Confirming Delivery Schedules & Automated Customer Notification", level=2)
    add_bullet("In the Order Details page, locate the 'Fulfillment & Dispatch Schedule' card.", bold_prefix="Step 1: ")
    add_bullet("Select the Confirmed Delivery Date (e.g. 2026-10-15) and Confirmed Time Window (e.g. 10:00 AM - 01:00 PM).", bold_prefix="Step 2: ")
    add_bullet("Add internal driver or cold-truck dispatch notes (optional).", bold_prefix="Step 3: ")
    add_bullet("Click 'Save Schedule & Notify Customer'. The platform automatically dispatches an official branded schedule confirmation email to the customer.", bold_prefix="Step 4 — Notify: ")

    add_heading("7.3. Printing Formal Commercial Tax Invoices", level=2)
    add_bullet("In any order details page, click the 'Print Invoice' button at the top right.", bold_prefix="Step 1: ")
    add_bullet("The system opens a clean, professionally styled commercial tax invoice ready for printing or saving as PDF.", bold_prefix="Step 2: ")
    add_bullet("The invoice includes MST company header, buyer SSM registration details, itemized SKU breakdown, delivery fee line item, and the grand total written in formal Ringgit words.", bold_prefix="Step 3: ")

    add_image_figure(os.path.join(assets_dir, "live_order_confirmed_24.png"), "Live Order Summary with Dispatch & Fulfillment Scheduling Details")

    # ═══════════════════════════════════════════════════════════════════════════
    # 9. CHAPTER 8: CUSTOMER VETTING & RFQ QUOTATIONS
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("8. Customer Vetting (Wholesale/Trading) & RFQ Quotations", level=1)

    add_heading("8.1. Approving B2B Wholesale Applicants", level=2)
    add_bullet("Navigate to: Admin Panel → Customer Accounts (/admin/customers).", bold_prefix="Step 1: ")
    add_bullet("Look for accounts with the yellow badge 'Pending Approval'.", bold_prefix="Step 2: ")
    add_bullet("Review the business details submitted by the applicant: Company Legal Name, SSM Business Registration Number, Contact Person, Phone, and Business Address.", bold_prefix="Step 3 — Vetting: ")
    add_bullet("Duplicate Alert: The system will alert you in red if the phone number or SSM number is already registered to prevent duplicate trade accounts.", bold_prefix="Step 4 — Alerts: ")
    add_bullet("Actions: Click 'Approve Account' to activate wholesale pricing access, or 'Reject' to decline. The applicant is automatically notified via transactional email.", bold_prefix="Step 5 — Action: ")

    add_heading("8.2. Managing Trading RFQ Quotations", level=2)
    add_bullet("Navigate to: Admin Panel → Quotations Management (/admin/quotations).", bold_prefix="Step 1: ")
    add_bullet("Click on any pending quotation to review requested bulk seafood quantities.", bold_prefix="Step 2: ")
    add_bullet("Enter the negotiated price per unit, add logistics notes, and set the Quote Expiry Date (validity window).", bold_prefix="Step 3: ")
    add_bullet("Click 'Send Formal Quote'. The client can now convert this quotation into an active order with 1 click from their account dashboard.", bold_prefix="Step 4: ")

    # ═══════════════════════════════════════════════════════════════════════════
    # 10. CHAPTER 9: DATABASE BACKUPS & SECURITY
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("9. 1-Click Database Backups & System Maintenance", level=1)
    add_body(
        "MST has complete ownership of its data. You can download a complete backup of your entire database at any time."
    )

    add_bullet("Navigate to: Admin Panel → Store Settings → Database Management tab (/admin/settings?tab=database).", bold_prefix="Step 1: ")
    add_bullet("Click the green button: 'Download Full Database Backup (.sql)'.", bold_prefix="Step 2: ")
    add_bullet("The system immediately dumps your complete MySQL database (including all products, orders, customers, settings, and invoices) and downloads it directly to your computer as a .sql file.", bold_prefix="Step 3 — Download: ")
    add_bullet("Recommended Routine: We recommend downloading a backup copy once a week and saving it to your secure company Google Drive or external drive.", bold_prefix="Step 4 — Best Practice: ")

    # ═══════════════════════════════════════════════════════════════════════════
    # 11. CHAPTER 10: QUICK TROUBLESHOOTING & FAQS
    # ═══════════════════════════════════════════════════════════════════════════
    add_heading("10. Troubleshooting & Frequently Asked Questions (FAQ)", level=1)

    add_bullet("Go to Admin Panel → Multilingual Translations (/admin/translations) and click the 'Clear Translation Cache' button at the top right. Then refresh your browser.", bold_prefix="Q1: I edited a translation or banner, but it doesn't show on the website immediately. ")
    add_bullet("Ensure that the product's 'Stock Quantity' is greater than 0, 'is_active' is checked, and if it's for walk-in shoppers, that 'is_walkin_available' is checked.", bold_prefix="Q2: A product appears as 'Out of Stock' or is hidden. ")
    add_bullet("Wholesale and Trading applicants cannot log in until approved. Go to Admin Panel → Customers, verify the account, and click 'Approve Account'.", bold_prefix="Q3: A wholesale customer says they cannot log in. ")
    add_bullet("Orders under RM150 residing in Zone A are automatically charged RM10. Check if the cart value is RM150 or above (which waives the fee to RM0). For outstation orders, the fee is intentionally shown as 'To Be Confirmed' with RM0 billed on Stripe.", bold_prefix="Q4: Why is a customer being charged RM10 for delivery? ")

    # Closing Signature Block
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    sig_p = doc.add_paragraph()
    sig_p.paragraph_format.space_before = Pt(10)
    sig_p.paragraph_format.space_after = Pt(2)
    sig_p.paragraph_format.line_spacing = 1.15
    r_sig1 = sig_p.add_run("Operations Manual Prepared By:\n\n")
    r_sig1.font.name = "Calibri"
    r_sig1.font.size = Pt(9.5)
    r_sig1.font.color.rgb = c_slate

    r_sig2 = sig_p.add_run("Abdul Rehman\n")
    r_sig2.font.name = "Calibri"
    r_sig2.font.bold = True
    r_sig2.font.size = Pt(10.5)
    r_sig2.font.color.rgb = c_navy

    r_sig3 = sig_p.add_run("Lead Full-Stack Developer & Technical Architect\nMST E-Commerce Platform Project")
    r_sig3.font.name = "Calibri"
    r_sig3.font.size = Pt(9)
    r_sig3.font.italic = True
    r_sig3.font.color.rgb = c_slate

    # Save Document
    output_filename = "f:\\My AI\\Sea Food\\seafood\\MST_Website_Admin_Panel_and_Operations_Manual.docx"
    doc.save(output_filename)
    print(f"Manual successfully created at: {output_filename}")

if __name__ == "__main__":
    create_manual()
