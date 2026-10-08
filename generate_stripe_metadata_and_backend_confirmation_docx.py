import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml import OxmlElement

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}" w:color="auto" w:val="clear"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_border(table, color="CBD5E1"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def add_heading(doc, text, level=1, color=None, space_before=12, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(12)
    elif level == 2:
        run.font.size = Pt(10.5)
    elif level == 3:
        run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = color
    return p

def add_body(doc, text, size=9.5, italic=False, bold=False, color=None, space_before=2, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.italic = italic
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    return p

def add_bullet(doc, text, size=9.5, color=None, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.font.name = "Calibri"
        r1.font.size = Pt(size)
        r1.font.bold = True
        if color:
            r1.font.color.rgb = color
        r2 = p.add_run(text)
        r2.font.name = "Calibri"
        r2.font.size = Pt(size)
        if color:
            r2.font.color.rgb = color
    else:
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p

def add_callout(doc, text, title=None, fill_hex="F0FDFA", border_hex="0F766E", text_color=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tblBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tblBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        r_title = p.add_run(f"{title}\n")
        r_title.font.name = "Calibri"
        r_title.font.size = Pt(9.5)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(15, 118, 110)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9)
    if text_color:
        r_text.font.color.rgb = text_color
    else:
        r_text.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(2)

def generate_document():
    doc = docx.Document()
    
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
    
    c_primary = RGBColor(15, 118, 110)    # Deep Teal
    c_dark = RGBColor(15, 23, 42)         # Slate 900
    c_sub = RGBColor(51, 65, 85)          # Slate 700
    c_gray = RGBColor(100, 116, 139)      # Slate 500
    
    # ─── HEADER BANNER ───────────────────────────────────────────────────────
    header_tbl = doc.add_table(rows=1, cols=1)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tbl.autofit = False
    
    h_cell = header_tbl.cell(0, 0)
    h_cell.width = Inches(6.5)
    set_cell_background(h_cell, "0F172A")
    set_cell_margins(h_cell, top=130, bottom=130, left=150, right=150)
    
    hp = h_cell.paragraphs[0]
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(2)
    
    r1 = hp.add_run("MST IMPORT & EXPORT SDN. BHD. — E-COMMERCE PLATFORM\n")
    r1.font.name = "Calibri"
    r1.font.size = Pt(9)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(204, 251, 241)
    
    r2 = hp.add_run("Stripe Test Key Receipt, Payment Metadata & Backend Storage Confirmation\n")
    r2.font.name = "Calibri"
    r2.font.size = Pt(13.5)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(255, 255, 255)
    
    r3 = hp.add_run("Publishable Key Linked · PaymentIntent Metadata Pipeline Fixed · Complete Order Persistence Confirmed")
    r3.font.name = "Calibri"
    r3.font.size = Pt(8.5)
    r3.font.italic = True
    r3.font.color.rgb = RGBColor(226, 232, 240)
    
    # Metadata Table
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    
    meta_tbl = doc.add_table(rows=2, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    set_table_border(meta_tbl, "E2E8F0")
    
    meta_data = [
        [("TO:", " Client Project Lead & MST Management Team"),
         ("DATE:", " October 8, 2026")],
        [("FROM:", " Abdul (Lead Technical Solution Architect)"),
         ("STATUS:", " Keys Linked · Stripe Payment Metadata Fixed · Backend Verified")]
    ]
    
    for r_idx, row in enumerate(meta_data):
        for c_idx, (label, val) in enumerate(row):
            cell = meta_tbl.cell(r_idx, c_idx)
            cell.width = Inches(3.25)
            set_cell_background(cell, "F8FAFC" if r_idx == 0 else "FFFFFF")
            set_cell_margins(cell, top=45, bottom=45, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            rl = p.add_run(label + " ")
            rl.font.name = "Calibri"
            rl.font.size = Pt(8.5)
            rl.font.bold = True
            rl.font.color.rgb = c_sub
            rv = p.add_run(val)
            rv.font.name = "Calibri"
            rv.font.size = Pt(8.5)
            rv.font.color.rgb = c_dark
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # ─── EXECUTIVE MESSAGE / RESPONSE ─────────────────────────────────────────
    add_heading(doc, "Executive Response & Clarification", level=1, color=c_primary)
    
    add_body(doc, "Dear Client,", size=9.5, bold=True)
    add_body(doc, "Thank you for providing the Stripe Test Mode Publishable Key (pk_test_...). We have successfully paired it with your Restricted Test Key (rk_test_...) and verified the integration in our sandbox environment.", size=9.5)
    add_body(doc, "We have thoroughly investigated your observations regarding the MYR 28.90 test payment and implemented comprehensive enhancements to address both Stripe Payment Metadata and backend data synchronization.", size=9.5)
    
    # ─── SECTION 1: KEY PAIRING SUMMARY ───────────────────────────────────────
    add_heading(doc, "1. Active Stripe Test Credentials Configuration", level=2, color=c_primary)
    
    key_tbl = doc.add_table(rows=4, cols=3)
    key_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    key_tbl.autofit = False
    set_table_border(key_tbl, "CBD5E1")
    
    headers = ["Key Parameter", "Key Prefix / Identifiers", "Integration & Scope Status"]
    for i, h in enumerate(headers):
        cell = key_tbl.cell(0, i)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = c_dark
        
    key_rows = [
        ("Publishable Key (Client Provided)", "pk_test_51UO5gZJuOivXc24TVh98dGVBhzQiy...EJ5z", "Installed in .env & Setting Model (Test Mode Active)"),
        ("Restricted Secret Key (Client Provided)", "rk_test_51UO5gZJuOivXc24TIzFAjL48Jpox5...wowY", "Verified with Checkout Sessions & PaymentIntents Permissions"),
        ("Environment Mode & Currency", "Mode: Test (Sandbox) · Currency: MYR (RM)", "Aligned across Web Storefront, Hosted Checkout & Database")
    ]
    
    col_widths = [Inches(2.2), Inches(2.3), Inches(2.0)]
    for r_idx, row in enumerate(key_rows):
        for c_idx, val in enumerate(row):
            cell = key_tbl.cell(r_idx + 1, c_idx)
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=45, bottom=45, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.color.rgb = c_sub
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # ─── SECTION 2: STRIPE PAYMENT METADATA EXPLANATION & FIX ─────────────────
    add_heading(doc, "2. Stripe Payment Details Metadata Section — Resolution", level=2, color=c_primary)
    
    add_body(doc, "Why Metadata was previously blank in the Stripe Dashboard Payment Details:", size=9.5, bold=True)
    add_bullet(doc, "In Stripe Official Hosted Checkout, the products, SKUs, and line item amounts appear automatically in the Checkout Summary on Stripe's hosted page.", size=9, bold_prefix="Checkout Summary vs Payment Details: ")
    add_bullet(doc, "However, under the Stripe Dashboard's 'Payment details' view, Stripe inspects the underlying PaymentIntent object (pi_...). By default, top-level metadata attached to a Checkout Session does NOT automatically transfer to the PaymentIntent unless explicitly passed inside 'payment_intent_data.metadata'.", size=9, bold_prefix="PaymentIntent Metadata Pipeline: ")
    add_bullet(doc, "We have updated the checkout controller (CheckoutController.php) to generate an Order Number upfront and pass complete order, customer, and fulfillment parameters into both the Checkout Session metadata and payment_intent_data.metadata.", size=9, bold_prefix="Implementation Update: ")
    
    add_callout(doc, 
        "Every future Stripe test payment will now display rich custom metadata directly inside the Stripe Dashboard under 'Payment details → Metadata', complete with Order Number, Fulfillment Type, Delivery Schedule / Collection Slot, Full Address, and Customer Contact Details.",
        title="Resolution Summary"
    )
    
    # Metadata Fields Table
    add_body(doc, "Table of Metadata Fields Sent to Stripe on Every Checkout:", size=9.5, bold=True)
    
    meta_fields_tbl = doc.add_table(rows=8, cols=3)
    meta_fields_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_fields_tbl.autofit = False
    set_table_border(meta_fields_tbl, "CBD5E1")
    
    mf_headers = ["Metadata Key", "Example Value Sent to Stripe", "Business & Operational Purpose"]
    for i, h in enumerate(mf_headers):
        cell = meta_fields_tbl.cell(0, i)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = c_dark
        
    mf_rows = [
        ("order_number", "ORD-6705A1BC89", "Matches exact Order Number recorded in MST website backend"),
        ("fulfillment_type", "delivery OR self_collection", "Identifies shipping dispatch vs walk-in/counter pickup"),
        ("delivery_date / collection_date", "2026-10-10 OR Today", "Preferred fulfillment date selected by customer"),
        ("delivery_address / collection_point", "7 Jalan SILC 2/18, Iskandar Puteri, Johor", "Complete street delivery address OR Counter 2 pickup location"),
        ("collection_time", "08:30 AM - 10:30 AM", "Specific pickup time window (for self-collection orders)"),
        ("customer_name & phone", "John Tan | +60123456789", "Direct customer contact information for driver/warehouse"),
        ("shipping_fee_myr & expected_total", "15.00 | 128.50", "Clear financial reconciliation against Stripe charge amount")
    ]
    
    mf_widths = [Inches(2.2), Inches(2.2), Inches(2.1)]
    for r_idx, row in enumerate(mf_rows):
        for c_idx, val in enumerate(row):
            cell = meta_fields_tbl.cell(r_idx + 1, c_idx)
            cell.width = mf_widths[c_idx]
            set_cell_background(cell, "FFFFFF" if r_idx % 2 == 0 else "F8FAFC")
            set_cell_margins(cell, top=45, bottom=45, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.color.rgb = c_sub
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # ─── SECTION 3: MST BACKEND ORDER STORAGE CONFIRMATION ────────────────────
    add_heading(doc, "3. MST Website Backend Order Storage — Confirmation", level=2, color=c_primary)
    
    add_body(doc, "We confirm that the MST website database and Admin Order Management portal record and preserve all order and fulfillment information permanently:", size=9.5)
    
    add_bullet(doc, "Generated automatically and synchronized with Stripe metadata (e.g. ORD-XXXX) and counter collection token (e.g. W-001).", size=9, bold_prefix="Order Number & Token: ")
    add_bullet(doc, "Complete breakdown of ordered seafood items, SKU codes, quantities, price tiers (Retail, Wholesale, Trading), unit prices, and line subtotals.", size=9, bold_prefix="Product Item Details: ")
    add_bullet(doc, "Full shipping address (Street, City, State, Postcode), delivery date, zone calculations, and transportation fees for Delivery orders; collection date and time slot for Self-Collection orders.", size=9, bold_prefix="Fulfillment Information: ")
    add_bullet(doc, "Payment status ('paid'), payment method ('stripe'), Stripe Payment Reference ID (pi_... / cs_...), and paid timestamp.", size=9, bold_prefix="Payment Audit & Reconciliation: ")
    add_bullet(doc, "Customer name, email address, mobile phone number, and optional customer remarks/notes.", size=9, bold_prefix="Customer Profile: ")
    add_bullet(doc, "All records are accessible under Admin Panel → Orders & Deliveries (/admin/orders/{order}) with order status workflow management, dispatch tracking, customer email notifications, and printable PDF invoices/slips.", size=9, bold_prefix="Admin Portal Management: ")
    
    # ─── SECTION 4: INVITATION FOR TEST VERIFICATION ──────────────────────────
    add_heading(doc, "4. Next Steps & UAT Verification", level=2, color=c_primary)
    
    add_body(doc, "You are invited to execute another test payment on the storefront using Stripe Test Mode. Upon completion:", size=9.5)
    add_bullet(doc, "Check the Stripe Dashboard under 'Payments' → click into the payment.", size=9, bold_prefix="Step 1: ")
    add_bullet(doc, "Verify that the 'Metadata' section now displays the Order Number, Fulfillment Type, Delivery Address / Collection Slot, and Customer details.", size=9, bold_prefix="Step 2: ")
    add_bullet(doc, "Log into the MST Admin Portal (/admin/orders) to see the corresponding order record with matching order number, itemized lines, and payment reference.", size=9, bold_prefix="Step 3: ")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    add_body(doc, "Thank you for your collaboration and guidance as we finalize the platform for full User Acceptance Testing (UAT). Please let us know if any further adjustments are desired.", size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    add_body(doc, "Warm regards,", size=9.5)
    add_body(doc, "Abdul & the Technical Development Team\nMST Import & Export Sdn. Bhd. E-Commerce Integration", size=9.5, bold=True, color=c_primary)
    
    output_filename = "MST_Stripe_Payment_Metadata_and_Backend_Storage_Confirmation.docx"
    output_path = f"f:\\My AI\\Sea Food\\seafood\\{output_filename}"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    generate_document()
